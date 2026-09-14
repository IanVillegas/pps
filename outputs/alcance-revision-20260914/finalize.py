"""Rename worksheet metadata, a capability missing from the documented JS API.

Workbook cells and formatting are authored by artifact-tool. This package pass
only renames sheets and their formula references; all other ZIP parts pass through.
"""
from pathlib import Path
from zipfile import ZipFile
from io import BytesIO
import json
from lxml import etree as ET
from openpyxl import load_workbook

base = Path(__file__).parent
path = base / 'Alcance_avance_intermedio_actualizado.xlsx'
mapping = dict(json.loads((base / 'renames.json').read_text(encoding='utf-8')))
with ZipFile(path) as src:
    result = BytesIO()
    with ZipFile(result, 'w') as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename.endswith('.xml') and info.filename.startswith('xl/'):
                root = ET.fromstring(data)
                changed = False
                for node in root.iter():
                    tag = ET.QName(node).localname
                    if tag == 'sheet' and node.get('name') in mapping:
                        node.set('name', mapping[node.get('name')])
                        changed = True
                    if tag in ('f', 'definedName', 'formula', 'formula1', 'formula2') and node.text:
                        old = node.text
                        for a, b in mapping.items():
                            node.text = node.text.replace("'" + a + "'!", "'" + b + "'!")
                        changed |= node.text != old
                if changed:
                    data = ET.tostring(root, encoding='UTF-8', xml_declaration=True)
            dst.writestr(info, data)
path.write_bytes(result.getvalue())

original = load_workbook('C:/Ian/UNA/3 - PPS/7 - Septiembre 05/Alcance_avance_intermedio_control.xlsx', read_only=True, data_only=True)
updated = load_workbook(path, read_only=True, data_only=True)
formulas = load_workbook(path, read_only=True, data_only=False)
assert len(updated.sheetnames) == 10
for index, last in [(1,13),(2,21),(3,14)]:
    old = list(original.worksheets[index].iter_rows(min_row=2,max_row=last,max_col=5,values_only=True))
    new = list(updated.worksheets[index].iter_rows(min_row=2,max_row=last,max_col=5,values_only=True))
    assert old == new, ('historical mismatch', index)
for row in range(25,32):
    for col in range(1,6):
        assert original['Avance Intermedio'].cell(row,col).value == updated['Avance Intermedio'].cell(row,col).value
assert updated['General']['C10'].value == 310
assert updated['General']['G17'].value == 190
assert updated['General']['G19'].value == 333
assert updated['Avance Intermedio']['B6'].value == 143
assert 'Plan previo' in updated['General']['D4'].value
errors=[]
for sheet in formulas:
    for row in sheet.iter_rows(max_row=60,max_col=25):
        for cell in row:
            if cell.data_type == 'e': errors.append((sheet.title,cell.coordinate,cell.value))
            if cell.data_type == 'f':
                assert not any("'"+old+"'!" in cell.value for old in mapping), cell.coordinate
assert not errors, errors
with ZipFile(path) as z:
    charts=[n for n in z.namelist() if '/charts/chart' in n and n.endswith('.xml')]
    assert len(charts)==4
print(json.dumps({'file':str(path),'sheets':updated.sheetnames,'futureHours':190,'actualHoursPreserved':143,'planHours':310,'forecastTotalHours':333,'formulaErrors':errors,'charts':len(charts),'historicalCellsPreserved':True},ensure_ascii=True))
